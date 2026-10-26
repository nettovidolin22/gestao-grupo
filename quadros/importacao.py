"""Importa um quadro exportado do Monday (Excel) ou uma planilha CSV simples.

O export do Monday tem esta cara:

    Nome do quadro
    Nome do grupo 1
    Name | Person | Status | Date
    Item A | Ana | Done | 2026-11-30
    (linha em branco)
    Nome do grupo 2
    Name | Person | Status | Date
    ...

Cada grupo vira um Grupo, e o tipo de cada coluna é adivinhado pelos valores.
"""

import csv
import io
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from itertools import cycle

from django.core.exceptions import ValidationError
from django.db import transaction

from .models import CORES, Coluna, Quadro

TAMANHO_MAXIMO = 5 * 1024 * 1024
LINHAS_MAXIMAS = 5000
COLUNAS_MAXIMAS = 100
# Um .xlsx é um zip: limita o tamanho descompactado para barrar "bombas de zip".
DESCOMPACTADO_MAXIMO = 50 * 1024 * 1024
CABECALHOS_NOME = {"name", "nome", "item", "itens", "tarefa", "tarefas"}
COLUNAS_IGNORADAS = {"subitems", "subitens", "item id", "id do item"}
NOMES_STATUS = {"status", "prioridade", "priority", "estado", "situação", "situacao", "fase"}
CORES_CONHECIDAS = {
    "done": "verde", "feito": "verde", "concluído": "verde", "concluido": "verde", "pronto": "verde",
    "working on it": "laranja", "fazendo": "laranja", "em andamento": "laranja",
    "stuck": "vermelho", "parado": "vermelho", "travado": "vermelho",
    "high": "vermelho", "alta": "vermelho", "critical": "vermelho", "crítica": "vermelho",
    "medium": "amarelo", "média": "amarelo", "media": "amarelo",
    "low": "azul", "baixa": "azul",
    "not started": "cinza", "a fazer": "cinza", "não iniciado": "cinza",
}


class ErroImportacao(Exception):
    pass


def ler_linhas(arquivo):
    """Devolve a planilha como lista de linhas, cada uma uma lista de valores."""
    if arquivo.size > TAMANHO_MAXIMO:
        raise ErroImportacao("O arquivo passa de 5 MB.")
    nome = arquivo.name.lower()
    if nome.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            with zipfile.ZipFile(arquivo) as pacote:
                descompactado = sum(info.file_size for info in pacote.infolist())
        except zipfile.BadZipFile:
            raise ErroImportacao("Não consegui abrir o Excel. Exporte de novo no formato .xlsx.")
        if descompactado > DESCOMPACTADO_MAXIMO:
            raise ErroImportacao("O Excel é grande demais por dentro. Divida a planilha em partes.")
        arquivo.seek(0)
        try:
            planilha = load_workbook(arquivo, read_only=True, data_only=True).worksheets[0]
            linhas = []
            for linha in planilha.iter_rows(values_only=True, max_col=COLUNAS_MAXIMAS):
                linhas.append(list(linha))
                if len(linhas) > LINHAS_MAXIMAS:
                    break
        except Exception:
            raise ErroImportacao("Não consegui abrir o Excel. Exporte de novo no formato .xlsx.")
    elif nome.endswith(".csv"):
        bruto = arquivo.read()
        for codificacao in ("utf-8-sig", "latin-1"):
            try:
                texto = bruto.decode(codificacao)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialeto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t")
        except csv.Error:
            dialeto = csv.excel
        linhas = []
        try:
            for linha in csv.reader(io.StringIO(texto), dialeto):
                linhas.append(linha[:COLUNAS_MAXIMAS])
                if len(linhas) > LINHAS_MAXIMAS:
                    break
        except csv.Error:
            raise ErroImportacao("Não consegui ler o CSV. Confira se o arquivo não está corrompido.")
    else:
        raise ErroImportacao("Envie um arquivo .xlsx ou .csv.")
    if len(linhas) > LINHAS_MAXIMAS:
        raise ErroImportacao(f"A planilha tem mais de {LINHAS_MAXIMAS} linhas.")
    return [[limpar_celula(c) for c in linha] for linha in linhas]


def limpar_celula(valor):
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def preenchidas(linha):
    return [c for c in linha if c != ""]


def eh_cabecalho(linha):
    valores = preenchidas(linha)
    return (
        len(valores) >= 2
        and linha[0] != ""
        and str(linha[0]).strip().lower() in CABECALHOS_NOME
    )


def separar_grupos(linhas):
    """Encontra nome do quadro, grupos, cabeçalhos e itens."""
    nome_quadro, grupos, titulos = "", [], []
    i = 0
    while i < len(linhas):
        linha = linhas[i]
        if eh_cabecalho(linha):
            if not grupos and titulos:
                nome_quadro = str(titulos[0])
            nome_grupo = str(titulos[-1]) if titulos else "Itens"
            cabecalho = [str(c).strip() for c in linha]
            itens = []
            i += 1
            # Itens seguem até uma linha em branco ou um novo cabeçalho.
            while i < len(linhas) and preenchidas(linhas[i]) and not eh_cabecalho(linhas[i]):
                if linhas[i][0] != "":  # linhas com a 1ª célula vazia são subitens: ignoradas
                    itens.append(linhas[i])
                i += 1
            grupos.append({"nome": nome_grupo, "cabecalho": cabecalho, "itens": itens})
            titulos = []
            continue
        valores = preenchidas(linha)
        if len(valores) == 1:
            titulos.append(valores[0])
        i += 1
    if not grupos:
        grupos = planilha_simples(linhas)
    if not grupos or not any(g["itens"] for g in grupos):
        raise ErroImportacao("Não encontrei itens na planilha. A primeira coluna deve se chamar Name ou Nome.")
    return nome_quadro, grupos


def planilha_simples(linhas):
    """CSV sem estrutura do Monday: primeira linha preenchida é o cabeçalho."""
    linhas = [l for l in linhas if preenchidas(l)]
    if len(linhas) < 2:
        return []
    cabecalho = [str(c).strip() for c in linhas[0]]
    return [{"nome": "Itens", "cabecalho": cabecalho, "itens": [l for l in linhas[1:] if l and l[0] != ""]}]


def como_data(valor):
    if isinstance(valor, date):
        return valor
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(str(valor), formato).date()
        except ValueError:
            continue
    return None


def como_numero(valor):
    if isinstance(valor, date):
        return None
    try:
        return Decimal(str(valor).replace(".", "").replace(",", ".") if "," in str(valor) else str(valor))
    except InvalidOperation:
        return None


def mapa_pessoas(filial):
    quadro_temporario = Quadro(filial=filial)
    mapa = {}
    for pessoa in quadro_temporario.pessoas():
        mapa[pessoa.username.lower()] = pessoa.pk
        if pessoa.get_full_name():
            mapa[pessoa.get_full_name().lower()] = pessoa.pk
    return mapa


def adivinhar_tipo(nome, valores, pessoas):
    if not valores:
        return Coluna.Tipo.TEXTO, []
    if all(como_data(v) for v in valores):
        return Coluna.Tipo.DATA, []
    if all(como_numero(v) is not None for v in valores):
        return Coluna.Tipo.NUMERO, []
    if all(str(v).lower() in pessoas for v in valores):
        return Coluna.Tipo.PESSOA, []
    distintos = list(dict.fromkeys(str(v) for v in valores))
    if len(distintos) <= 12 and (nome.lower() in NOMES_STATUS or len(valores) > len(distintos)):
        cores = cycle(c for c in CORES if c != "cinza")
        opcoes = [
            {"nome": d[:40], "cor": CORES_CONHECIDAS.get(d.lower()) or next(cores)} for d in distintos
        ]
        return Coluna.Tipo.STATUS, opcoes
    return Coluna.Tipo.TEXTO, []


def valor_para_coluna(coluna, valor, pessoas):
    if coluna.tipo == Coluna.Tipo.DATA:
        return como_data(valor).isoformat()
    if coluna.tipo == Coluna.Tipo.NUMERO:
        return str(como_numero(valor))
    if coluna.tipo == Coluna.Tipo.PESSOA:
        return str(pessoas[str(valor).lower()])
    return str(valor)


def importar_planilha(arquivo, nome, filial, usuario):
    linhas = ler_linhas(arquivo)
    nome_arquivo, grupos = separar_grupos(linhas)
    pessoas = mapa_pessoas(filial)

    # Colunas na ordem em que aparecem, juntando os cabeçalhos de todos os grupos.
    nomes_colunas = []
    for grupo in grupos:
        for titulo in grupo["cabecalho"][1:]:
            if titulo and titulo.lower() not in COLUNAS_IGNORADAS and titulo not in nomes_colunas:
                nomes_colunas.append(titulo)
    valores_por_coluna = {n: [] for n in nomes_colunas}
    for grupo in grupos:
        for linha in grupo["itens"]:
            for posicao, titulo in enumerate(grupo["cabecalho"]):
                if titulo in valores_por_coluna and posicao < len(linha) and linha[posicao] != "":
                    valores_por_coluna[titulo].append(linha[posicao])

    with transaction.atomic():
        quadro = Quadro.objects.create(
            filial=filial,
            nome=(nome or nome_arquivo or "Quadro importado")[:120],
            descricao=f"Importado de {arquivo.name}",
            criado_por=usuario,
        )
        colunas = {}
        for ordem, titulo in enumerate(nomes_colunas):
            tipo, opcoes = adivinhar_tipo(titulo, valores_por_coluna[titulo], pessoas)
            colunas[titulo] = quadro.colunas.create(nome=titulo[:60], tipo=tipo, opcoes=opcoes, ordem=ordem)
        for ordem_grupo, dados in enumerate(grupos):
            cor = CORES[(ordem_grupo + 1) % len(CORES)]
            grupo = quadro.grupos.create(nome=dados["nome"][:120], cor=cor, ordem=ordem_grupo)
            for ordem_item, linha in enumerate(dados["itens"]):
                valores = {}
                for posicao, titulo in enumerate(dados["cabecalho"]):
                    coluna = colunas.get(titulo)
                    if coluna and posicao < len(linha) and linha[posicao] != "":
                        try:
                            valor = coluna.limpar_valor(valor_para_coluna(coluna, linha[posicao], pessoas))
                        except (ValidationError, KeyError, AttributeError):
                            continue
                        valores[str(coluna.pk)] = valor
                grupo.itens.create(
                    nome=str(linha[0])[:255], ordem=ordem_item, valores=valores, criado_por=usuario
                )
    return quadro
